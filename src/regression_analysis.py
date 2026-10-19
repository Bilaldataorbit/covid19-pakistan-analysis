"""
src/regression_analysis.py
==========================
The regression fundamentals module: every model the course asks for, written so the
notebook can call one function per topic and the report can export the same numbers.

Covered
-------
1. Simple linear regression with statsmodels  -> sm.OLS(...).summary()
2. Simple linear regression with scikit-learn -> LinearRegression on Day_Num
3. Multiple linear regression + VIF            -> Deaths, Recovered, Day_Num
4. Regression diagnostics                      -> Durbin-Watson, Breusch-Pagan, Jarque-Bera
5. Polynomial regression, degree comparison    -> degrees 1 to 5
6. Logistic regression on High_Cases           -> sm.Logit
7. ROC curve, AUC, confusion matrix            -> classification metrics
8. Cross-validation with Ridge and Lasso       -> KFold over alpha grids
9. Excel export                                -> Regression_Analysis_Report.xlsx

The target throughout is daily cases; the predictors are Day_Num, deaths and recoveries
from the same day, which is what a regression *fundamentals* exercise asks for. This is a
descriptive model of the epidemic curve, deliberately different from the leakage-free
forecasting model in ml_models.py (that one uses lagged predictors and a time split).
"""
from __future__ import annotations

import os
from dataclasses import dataclass, field

import numpy as np
import pandas as pd
import statsmodels.api as sm
from scipy import stats
from sklearn.linear_model import Lasso, LinearRegression, LogisticRegression, Ridge
from sklearn.metrics import (accuracy_score, auc, classification_report,
                             confusion_matrix, f1_score, precision_score,
                             recall_score, roc_auc_score, roc_curve)
from sklearn.model_selection import KFold, cross_val_score
from sklearn.preprocessing import PolynomialFeatures, StandardScaler
from statsmodels.stats.diagnostic import het_breuschpagan
from statsmodels.stats.outliers_influence import variance_inflation_factor
from statsmodels.stats.stattools import durbin_watson, jarque_bera

from data_pipeline import CLEAN_PATH, ROOT

OUT_DIR = os.path.join(ROOT, "outputs")
os.makedirs(OUT_DIR, exist_ok=True)

TARGET = "Cases"
HIGH_FLAG = "High_Cases"
SEED = 42


# --------------------------------------------------------------------- data
def load_regression_data(path: str | None = None) -> pd.DataFrame:
    """One row per observed calendar date (the 8 gap days are excluded so a
    Day_Num regressor is not fitted against artificial zeros)."""
    df = pd.read_csv(path or CLEAN_PATH, parse_dates=["Date"])
    d = (df.groupby("Date")
           .agg(Cases=("Cases", "sum"), Deaths=("Deaths", "sum"),
                Recovered=("Recovered", "sum"))
           .reset_index()
           .sort_values("Date")
           .reset_index(drop=True))
    d["Day_Num"] = np.arange(1, len(d) + 1)
    d["Day_Squared"] = d["Day_Num"] ** 2
    d["High_Cases"] = (d["Cases"] > 100).astype(int)
    d["Cumulative_Cases"] = d["Cases"].cumsum()
    return d


# ------------------------------------------------------- 1. statsmodels OLS
def simple_lr_statsmodels(d: pd.DataFrame | None = None):
    """Cases = b0 + b1 * Day_Num, fitted with sm.OLS so we get the full summary."""
    d = load_regression_data() if d is None else d
    X = sm.add_constant(d[["Day_Num"]].astype(float))
    y = d[TARGET].astype(float)
    model = sm.OLS(y, X).fit()

    out = dict(
        model=model,
        summary=str(model.summary()),
        params=model.params.to_dict(),
        pvalues=model.pvalues.to_dict(),
        conf_int=model.conf_int().to_dict(),
        r2=float(model.rsquared),
        adj_r2=float(model.rsquared_adj),
        fvalue=float(model.fvalue),
        f_pvalue=float(model.f_pvalue),
        aic=float(model.aic),
        bic=float(model.bic),
        equation=("Cases = %.2f %+ .4f * Day_Num" % (model.params["const"], model.params["Day_Num"])),
    )
    out["table"] = pd.DataFrame({
        "term": model.params.index,
        "coefficient": model.params.values,
        "std_error": model.bse.values,
        "t_stat": model.tvalues.values,
        "p_value": model.pvalues.values,
        "ci_lower": model.conf_int()[0].values,
        "ci_upper": model.conf_int()[1].values,
    })
    return out


# -------------------------------------------------------- 2. sklearn simple
def simple_lr_sklearn(d: pd.DataFrame | None = None):
    """Same model through scikit-learn, so the coefficients can be compared."""
    d = load_regression_data() if d is None else d
    X = d[["Day_Num"]].to_numpy(float)
    y = d[TARGET].to_numpy(float)
    lr = LinearRegression().fit(X, y)
    pred = lr.predict(X)
    resid = y - pred
    return dict(
        model=lr,
        intercept=float(lr.intercept_),
        slope=float(lr.coef_[0]),
        r2=float(lr.score(X, y)),
        rmse=float(np.sqrt(np.mean(resid ** 2))),
        mae=float(np.mean(np.abs(resid))),
        predictions=pred,
        residuals=resid,
        equation="Cases = %.2f %+ .4f * Day_Num" % (lr.intercept_, lr.coef_[0]),
    )


# ------------------------------------------------- 3. multiple regression VIF
def multiple_lr_vif(d: pd.DataFrame | None = None):
    """Cases = b0 + b1*Deaths + b2*Recovered + b3*Day_Num, plus a VIF table."""
    d = load_regression_data() if d is None else d
    predictors = ["Deaths", "Recovered", "Day_Num"]
    X = sm.add_constant(d[predictors].astype(float))
    y = d[TARGET].astype(float)
    model = sm.OLS(y, X).fit()

    vif_rows = []
    for i, col in enumerate(X.columns):
        if col == "const":
            continue
        vif_rows.append({
            "feature": col,
            "VIF": float(variance_inflation_factor(X.values, i)),
            # tolerance = 1/VIF, the textbook companion number
            "tolerance": float(1.0 / variance_inflation_factor(X.values, i)),
        })
    vif = pd.DataFrame(vif_rows).sort_values("VIF", ascending=False).reset_index(drop=True)
    vif["verdict"] = np.where(vif["VIF"] < 5, "OK",
                              np.where(vif["VIF"] < 10, "Moderate", "Severe"))

    corr = d[predictors].corr().round(3)

    table = pd.DataFrame({
        "term": model.params.index,
        "coefficient": model.params.values,
        "std_error": model.bse.values,
        "t_stat": model.tvalues.values,
        "p_value": model.pvalues.values,
        "ci_lower": model.conf_int()[0].values,
        "ci_upper": model.conf_int()[1].values,
    })
    return dict(
        model=model, table=table, vif=vif, correlation=corr,
        summary=str(model.summary()),
        r2=float(model.rsquared), adj_r2=float(model.rsquared_adj),
        fvalue=float(model.fvalue), f_pvalue=float(model.f_pvalue),
        equation=("Cases = %.2f %+ .4f*Deaths %+ .4f*Recovered %+ .4f*Day_Num"
                  % tuple(model.params.reindex(["const"] + predictors))),
    )


# ------------------------------------------------------- 4. diagnostics
def regression_diagnostics(d: pd.DataFrame | None = None):
    """Durbin-Watson, Breusch-Pagan and Jarque-Bera on the multiple model."""
    d = load_regression_data() if d is None else d
    res = multiple_lr_vif(d)
    model = res["model"]

    # Durbin-Watson: ~2 means no first-order autocorrelation
    dw = float(durbin_watson(model.resid))

    # Breusch-Pagan: H0 = homoscedastic residuals
    bp_stat, bp_p, bp_f, bp_fp = het_breuschpagan(model.resid, model.model.exog)

    # Jarque-Bera: H0 = residuals are normal
    jb_stat, jb_p, jb_skew, jb_kurt = jarque_bera(model.resid)

    # Shapiro-Wilk as a second normality opinion (n <= 5000 here)
    sw_stat, sw_p = stats.shapiro(model.resid)

    resid = pd.Series(model.resid, name="residual")
    return dict(
        durbin_watson=dw,
        dw_verdict=("no serious autocorrelation" if 1.5 < dw < 2.5
                    else "positive autocorrelation" if dw <= 1.5
                    else "negative autocorrelation"),
        breusch_pagan_stat=float(bp_stat),
        breusch_pagan_p=float(bp_p),
        breusch_pagan_verdict=("homoscedastic" if bp_p > 0.05 else "heteroscedastic"),
        jarque_bera_stat=float(jb_stat),
        jarque_bera_p=float(jb_p),
        jarque_bera_skew=float(jb_skew),
        jarque_bera_kurtosis=float(jb_kurt),
        jarque_bera_verdict=("residuals look normal" if jb_p > 0.05 else "residuals are not normal"),
        shapiro_stat=float(sw_stat),
        shapiro_p=float(sw_p),
        residuals=resid,
        fitted=pd.Series(model.fittedvalues, name="fitted"),
        summary_table=pd.DataFrame([
            {"test": "Durbin-Watson", "statistic": dw, "p_value": np.nan,
             "null_hypothesis": "no autocorrelation", "verdict": "no serious autocorrelation" if 1.5 < dw < 2.5 else "autocorrelation present"},
            {"test": "Breusch-Pagan", "statistic": float(bp_stat), "p_value": float(bp_p),
             "null_hypothesis": "homoscedasticity", "verdict": "homoscedastic" if bp_p > 0.05 else "heteroscedastic"},
            {"test": "Jarque-Bera", "statistic": float(jb_stat), "p_value": float(jb_p),
             "null_hypothesis": "residual normality", "verdict": "normal" if jb_p > 0.05 else "not normal"},
            {"test": "Shapiro-Wilk", "statistic": float(sw_stat), "p_value": float(sw_p),
             "null_hypothesis": "residual normality", "verdict": "normal" if sw_p > 0.05 else "not normal"},
        ]),
    )


# --------------------------------------------------- 5. polynomial comparison
def polynomial_comparison(d: pd.DataFrame | None = None, max_degree: int = 5):
    """Fit degrees 1..max_degree and compare R2, adjusted R2, AIC and RMSE."""
    d = load_regression_data() if d is None else d
    y = d[TARGET].to_numpy(float)
    x = d[["Day_Num"]].to_numpy(float)
    n = len(y)

    rows, models = [], {}
    for deg in range(1, max_degree + 1):
        poly = PolynomialFeatures(degree=deg, include_bias=False)
        Xp = poly.fit_transform(x)
        Xc = sm.add_constant(Xp)
        m = sm.OLS(y, Xc).fit()
        pred = m.predict(Xc)
        rmse = float(np.sqrt(np.mean((y - pred) ** 2)))
        rows.append({
            "degree": deg,
            "r2": float(m.rsquared),
            "adj_r2": float(m.rsquared_adj),
            "aic": float(m.aic),
            "bic": float(m.bic),
            "rmse": rmse,
            "f_pvalue": float(m.f_pvalue),
        })
        models[deg] = dict(model=m, poly=poly, pred=pred, X=Xc)

    table = pd.DataFrame(rows)
    best_r2 = int(table.loc[table["r2"].idxmax(), "degree"])
    best_aic = int(table.loc[table["aic"].idxmin(), "degree"])
    best_adj = int(table.loc[table["adj_r2"].idxmax(), "degree"])
    return dict(table=table, models=models, best_degree_r2=best_r2,
                best_degree_aic=best_aic, best_degree_adj=best_adj,
                predictions={k: v["pred"] for k, v in models.items()})


# ------------------------------------------------------- 6. logistic regression
def logistic_high_cases(d: pd.DataFrame | None = None):
    """log(p/(1-p)) = b0 + b1*Day_Num + b2*Deaths for High_Cases (cases > 100)."""
    d = load_regression_data() if d is None else d
    X = sm.add_constant(d[["Day_Num", "Deaths"]].astype(float))
    y = d[HIGH_FLAG].astype(float)

    note = "fitted by maximum likelihood (sm.Logit)"
    try:
        model = sm.Logit(y, X).fit(disp=0)
    except Exception:
        # perfect separation is common on a monotone epidemic curve
        model = sm.Logit(y, X).fit_regularized(disp=0, alpha=0.01)
        note = "maximum likelihood did not converge (separation); fitted with L1-regularised Logit"

    proba = model.predict(X)
    pred = (proba > 0.5).astype(int)

    # odds ratios read better than raw log-odds
    odds = np.exp(model.params)

    table = pd.DataFrame({
        "term": model.params.index,
        "coefficient": model.params.values,
        "odds_ratio": odds.values,
        "std_error": model.bse.values,
        "z_stat": model.tvalues.values,
        "p_value": model.pvalues.values,
    })
    return dict(
        model=model, table=table, proba=proba, predicted=pred,
        note=note,
        pseudo_r2=float(getattr(model, "prsquared", np.nan)),
        llr_pvalue=float(getattr(model, "llr_pvalue", np.nan)),
        equation=("logit(p) = %.4f %+ .4f*Day_Num %+ .4f*Deaths"
                  % tuple(model.params.reindex(["const", "Day_Num", "Deaths"]))),
        summary=str(model.summary()),
    )


# ------------------------------------------- 7. ROC, AUC, confusion matrix
def roc_auc_confusion(d: pd.DataFrame | None = None):
    """Classification quality of the logistic model at the 0.5 cut-off, plus ROC."""
    d = load_regression_data() if d is None else d
    logit = logistic_high_cases(d)
    y = d[HIGH_FLAG].to_numpy(int)
    proba = np.asarray(logit["proba"])
    pred = (proba > 0.5).astype(int)

    cm = confusion_matrix(y, pred, labels=[0, 1])
    tn, fp, fn, tp = cm.ravel()
    fpr, tpr, thresholds = roc_curve(y, proba)
    auc_value = roc_auc_score(y, proba)

    # Youden J picks the best single threshold
    j = tpr - fpr
    best_i = int(np.argmax(j))

    metrics = pd.DataFrame([
        {"metric": "Accuracy", "value": accuracy_score(y, pred)},
        {"metric": "Precision", "value": precision_score(y, pred, zero_division=0)},
        {"metric": "Recall (sensitivity)", "value": recall_score(y, pred, zero_division=0)},
        {"metric": "F1 score", "value": f1_score(y, pred, zero_division=0)},
        {"metric": "Specificity", "value": tn / (tn + fp) if (tn + fp) else np.nan},
        {"metric": "False positive rate", "value": fp / (fp + tn) if (fp + tn) else np.nan},
        {"metric": "AUC (ROC)", "value": auc_value},
    ])

    cm_table = pd.DataFrame(
        [[int(tn), int(fp)], [int(fn), int(tp)]],
        index=["Actual 0 (<=100)", "Actual 1 (>100)"],
        columns=["Predicted 0", "Predicted 1"],
    )
    report = classification_report(y, pred, labels=[0, 1],
                                   target_names=["low", "high"], zero_division=0)

    return dict(
        confusion_matrix=cm_table, tn=int(tn), fp=int(fp), fn=int(fn), tp=int(tp),
        accuracy=float(accuracy_score(y, pred)),
        precision=float(precision_score(y, pred, zero_division=0)),
        recall=float(recall_score(y, pred, zero_division=0)),
        f1=float(f1_score(y, pred, zero_division=0)),
        specificity=float(tn / (tn + fp)) if (tn + fp) else np.nan,
        auc=float(auc_value),
        fpr=fpr, tpr=tpr, thresholds=thresholds,
        best_threshold=float(thresholds[best_i]),
        best_youden=float(j[best_i]),
        metrics=metrics, classification_report=report,
        logit=logit,
    )


# ---------------------------------- 8. cross-validation, Ridge and Lasso
def cross_validation_ridge_lasso(d: pd.DataFrame | None = None,
                                 predictors=None, folds: int = 5):
    """KFold cross-validation (not a time split: this is a fundamentals exercise).

    Compares OLS, Ridge and Lasso across an alpha grid on the multiple-regression
    feature set, reporting R2 with a mean and a standard deviation per model.
    """
    d = load_regression_data() if d is None else d
    predictors = predictors or ["Deaths", "Recovered", "Day_Num"]
    X = d[predictors].to_numpy(float)
    y = d[TARGET].to_numpy(float)
    kf = KFold(n_splits=folds, shuffle=True, random_state=SEED)

    rows = []

    ols_scores = cross_val_score(LinearRegression(), X, y, cv=kf, scoring="r2")
    rows.append({"model": "OLS", "alpha": np.nan,
                 "cv_r2_mean": ols_scores.mean(), "cv_r2_std": ols_scores.std()})

    for alpha in [0.01, 0.1, 1.0, 10.0, 100.0, 1000.0]:
        pipe = StandardScaler()
        Xs = pipe.fit_transform(X)
        ridge = cross_val_score(Ridge(alpha=alpha), Xs, y, cv=kf, scoring="r2")
        lasso = cross_val_score(Lasso(alpha=alpha, max_iter=50000), Xs, y, cv=kf, scoring="r2")
        rows.append({"model": "Ridge", "alpha": alpha,
                     "cv_r2_mean": ridge.mean(), "cv_r2_std": ridge.std()})
        rows.append({"model": "Lasso", "alpha": alpha,
                     "cv_r2_mean": lasso.mean(), "cv_r2_std": lasso.std()})

    table = pd.DataFrame(rows)
    best = table.loc[table["cv_r2_mean"].idxmax()]

    # refit the winners on all the data so the coefficients can be read
    Xs = StandardScaler().fit_transform(X)
    best_alpha = float(best["alpha"]) if not pd.isna(best["alpha"]) else 1.0
    ridge_full = Ridge(alpha=best_alpha).fit(Xs, y)
    lasso_full = Lasso(alpha=best_alpha, max_iter=50000).fit(Xs, y)

    coefs = pd.DataFrame({
        "feature": predictors,
        "ridge_coef": ridge_full.coef_,
        "lasso_coef": lasso_full.coef_,
    })
    return dict(table=table.sort_values("cv_r2_mean", ascending=False).reset_index(drop=True),
                best=best.to_dict(), folds=folds, predictors=predictors,
                coefficients=coefs, ridge=ridge_full, lasso=lasso_full,
                ols_scores=ols_scores)


# ------------------------------------------------------------- 9. run-all
@dataclass
class RegressionResults:
    data: pd.DataFrame
    simple_sm: dict = field(default_factory=dict)
    simple_sk: dict = field(default_factory=dict)
    multiple: dict = field(default_factory=dict)
    diagnostics: dict = field(default_factory=dict)
    polynomial: dict = field(default_factory=dict)
    logistic: dict = field(default_factory=dict)
    classification: dict = field(default_factory=dict)
    cv: dict = field(default_factory=dict)


def run_all(verbose: bool = True) -> RegressionResults:
    d = load_regression_data()
    res = RegressionResults(data=d)

    res.simple_sm = simple_lr_statsmodels(d)
    res.simple_sk = simple_lr_sklearn(d)
    res.multiple = multiple_lr_vif(d)
    res.diagnostics = regression_diagnostics(d)
    res.polynomial = polynomial_comparison(d)
    res.logistic = logistic_high_cases(d)
    res.classification = roc_auc_confusion(d)
    res.cv = cross_validation_ridge_lasso(d)

    if verbose:
        print("=" * 78)
        print("1. SIMPLE LINEAR REGRESSION - statsmodels OLS")
        print("=" * 78)
        print(res.simple_sm["summary"])
        print()
        print("=" * 78)
        print("2. SIMPLE LINEAR REGRESSION - scikit-learn")
        print("=" * 78)
        print("  equation   :", res.simple_sk["equation"])
        print("  R2         : %.4f" % res.simple_sk["r2"])
        print("  RMSE       : %.2f" % res.simple_sk["rmse"])
        print("  matches OLS:", np.isclose(res.simple_sk["slope"],
                                         res.simple_sm["params"]["Day_Num"]))
        print()
        print("=" * 78)
        print("3. MULTIPLE LINEAR REGRESSION + VIF")
        print("=" * 78)
        print(res.multiple["table"].round(4).to_string(index=False))
        print("\n  R2 = %.4f  adj R2 = %.4f  F = %.2f (p = %.3e)"
              % (res.multiple["r2"], res.multiple["adj_r2"],
                 res.multiple["fvalue"], res.multiple["f_pvalue"]))
        print("\n  VIF:")
        print(res.multiple["vif"].round(3).to_string(index=False))
        print()
        print("=" * 78)
        print("4. REGRESSION DIAGNOSTICS")
        print("=" * 78)
        print(res.diagnostics["summary_table"].round(4).to_string(index=False))
        print()
        print("=" * 78)
        print("5. POLYNOMIAL REGRESSION - degree comparison")
        print("=" * 78)
        print(res.polynomial["table"].round(4).to_string(index=False))
        print("\n  best by R2: degree %d | by adjusted R2: degree %d | by AIC: degree %d"
              % (res.polynomial["best_degree_r2"], res.polynomial["best_degree_adj"],
                 res.polynomial["best_degree_aic"]))
        print()
        print("=" * 78)
        print("6. LOGISTIC REGRESSION - High_Cases")
        print("=" * 78)
        print("  ", res.logistic["note"])
        print(res.logistic["table"].round(4).to_string(index=False))
        print("  pseudo R2 = %.4f" % res.logistic["pseudo_r2"])
        print()
        print("=" * 78)
        print("7. CONFUSION MATRIX / ROC / AUC")
        print("=" * 78)
        print(res.classification["confusion_matrix"].to_string())
        print()
        print(res.classification["metrics"].round(4).to_string(index=False))
        print("  best threshold by Youden J: %.2f" % res.classification["best_threshold"])
        print()
        print("=" * 78)
        print("8. CROSS-VALIDATION with Ridge and Lasso (%d folds)" % res.cv["folds"])
        print("=" * 78)
        print(res.cv["table"].round(4).to_string(index=False))
        print("\n  best:", res.cv["best"])

    return res




# ------------------------------------------------------- 9. Excel export
def export_report(res: RegressionResults | None = None, path: str | None = None) -> str:
    """Write Regression_Analysis_Report.xlsx with one sheet per requirement."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    res = res or run_all(verbose=False)
    path = path or os.path.join(OUT_DIR, "Regression_Analysis_Report.xlsx")

    NAVY = "1F3864"
    HDR = Font(name="Calibri", size=10, bold=True, color="FFFFFF")
    FILL = PatternFill("solid", fgColor=NAVY)
    TITLE = Font(name="Calibri", size=14, bold=True, color=NAVY)
    BODY = Font(name="Calibri", size=9)
    THIN = Side(style="thin", color="BFBFBF")
    BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, title, df=None, note=None, start=3):
        ws = wb.create_sheet(name[:31])
        ws.cell(row=1, column=1, value=title).font = TITLE
        if note:
            ws.cell(row=2, column=1, value=note).font = Font(name="Calibri", size=9,
                                                            italic=True, color="808080")
        if df is None:
            return ws
        r = start
        for j, col in enumerate(df.columns, 1):
            c = ws.cell(row=r, column=j, value=str(col))
            c.font = HDR; c.fill = FILL; c.border = BOX
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for i, row in enumerate(df.itertuples(index=False), start=r + 1):
            for j, v in enumerate(row, 1):
                if isinstance(v, (np.integer,)):
                    v = int(v)
                elif isinstance(v, (np.floating,)):
                    v = float(v) if not np.isnan(v) else ""
                c = ws.cell(row=i, column=j, value=v)
                c.border = BOX; c.font = BODY
        for j in range(1, len(df.columns) + 1):
            width = max(13, min(34, int(df[df.columns[j - 1]].astype(str).str.len().max() or 13) + 2))
            ws.column_dimensions[get_column_letter(j)].width = width
        ws.freeze_panes = ws.cell(row=r + 1, column=1)
        return ws

    # ---- overview
    ws = wb.create_sheet("Overview")
    ws["A1"] = "Regression Analysis Report - COVID-19 Pakistan"
    ws["A1"].font = Font(name="Calibri", size=17, bold=True, color=NAVY)
    ws["A2"] = ("Historical analysis, 26 Feb - 03 Jun 2020. The predictors are same-day "
                "Deaths / Recovered / Day_Num, which is what a regression fundamentals "
                "exercise asks for; it is descriptive, not a forecaster.")
    ws["A2"].font = Font(name="Calibri", size=9, italic=True, color="808080")
    ws["A2"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A2:F2")
    ws.row_dimensions[2].height = 28

    rows = [
        ("Dataset", ""),
        ("Observations", len(res.data)),
        ("Date range", f"{res.data.Date.min().date()} to {res.data.Date.max().date()}"),
        ("Target", TARGET),
        ("", ""),
        ("1. Simple LR - statsmodels OLS", ""),
        ("R squared", round(res.simple_sm["r2"], 6)),
        ("Adjusted R squared", round(res.simple_sm["adj_r2"], 6)),
        ("F statistic", round(res.simple_sm["fvalue"], 3)),
        ("Model p-value", res.simple_sm["f_pvalue"]),
        ("Equation", res.simple_sm["equation"]),
        ("", ""),
        ("2. Simple LR - scikit-learn", ""),
        ("Matches statsmodels", "yes" if np.isclose(res.simple_sk["slope"], res.simple_sm["params"]["Day_Num"]) else "no"),
        ("R squared", round(res.simple_sk["r2"], 6)),
        ("RMSE", round(res.simple_sk["rmse"], 3)),
        ("", ""),
        ("3. Multiple LR + VIF", ""),
        ("R squared", round(res.multiple["r2"], 6)),
        ("Adjusted R squared", round(res.multiple["adj_r2"], 6)),
        ("Max VIF", round(res.multiple["vif"]["VIF"].max(), 3)),
        ("Multicollinearity", res.multiple["vif"].iloc[0]["verdict"]),
        ("", ""),
        ("4. Diagnostics", ""),
        ("Durbin-Watson", round(res.diagnostics["durbin_watson"], 4)),
        ("Breusch-Pagan p", res.diagnostics["breusch_pagan_p"]),
        ("Jarque-Bera p", res.diagnostics["jarque_bera_p"]),
        ("", ""),
        ("5. Polynomial regression", ""),
        ("Best degree by adjusted R2", res.polynomial["best_degree_adj"]),
        ("Adjusted R2 at that degree", round(float(res.polynomial["table"].loc[
            res.polynomial["table"]["degree"] == res.polynomial["best_degree_adj"], "adj_r2"].iloc[0]), 6)),
        ("", ""),
        ("6. Logistic regression", ""),
        ("Pseudo R squared", round(res.logistic["pseudo_r2"], 6)),
        ("Converged by", res.logistic["note"]),
        ("", ""),
        ("7. Classification", ""),
        ("Confusion matrix", f"TN={res.classification['tn']} FP={res.classification['fp']} "
                             f"FN={res.classification['fn']} TP={res.classification['tp']}"),
        ("Accuracy", round(res.classification["accuracy"], 6)),
        ("Precision", round(res.classification["precision"], 6)),
        ("Recall", round(res.classification["recall"], 6)),
        ("F1", round(res.classification["f1"], 6)),
        ("AUC", round(res.classification["auc"], 6)),
        ("Best threshold (Youden J)", round(res.classification["best_threshold"], 4)),
        ("", ""),
        ("8. Cross-validation (5 folds)", ""),
        ("Best model", str(res.cv["best"]["model"])),
        ("Best alpha", res.cv["best"]["alpha"] if not pd.isna(res.cv["best"]["alpha"]) else "n/a"),
        ("CV R2 mean", round(float(res.cv["best"]["cv_r2_mean"]), 6)),
        ("CV R2 std", round(float(res.cv["best"]["cv_r2_std"]), 6)),
    ]
    r = 4
    for k, v in rows:
        if k and v == "":
            c = ws.cell(row=r, column=1, value=k)
            c.font = Font(name="Calibri", size=11, bold=True, color=NAVY)
            c.fill = PatternFill("solid", fgColor="DEEBF7")
            ws.cell(row=r, column=2).fill = PatternFill("solid", fgColor="DEEBF7")
        elif k:
            ws.cell(row=r, column=1, value=k).font = Font(name="Calibri", size=10, bold=True)
            c = ws.cell(row=r, column=2, value=v)
            c.font = BODY; c.alignment = Alignment(horizontal="left")
        r += 1
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 52

    # ---- per-requirement sheets
    sheet("1_Simple_LR_sm", "1. Simple linear regression - statsmodels OLS",
          res.simple_sm["table"], note=res.simple_sm["equation"])
    ws = wb["1_Simple_LR_sm"]
    ws.cell(row=res.simple_sm["table"].shape[0] + 6, column=1,
            value="Full sm.OLS summary").font = Font(bold=True, color=NAVY)
    for i, line in enumerate(res.simple_sm["summary"].splitlines(), 1):
        ws.cell(row=res.simple_sm["table"].shape[0] + 6 + i, column=1, value=line).font = Font(
            name="Consolas", size=8)

    sheet("2_Simple_LR_sklearn", "2. Simple linear regression - scikit-learn",
          pd.DataFrame([{
              "intercept": res.simple_sk["intercept"],
              "slope_Day_Num": res.simple_sk["slope"],
              "r2": res.simple_sk["r2"],
              "rmse": res.simple_sk["rmse"],
              "mae": res.simple_sk["mae"],
          }]), note=res.simple_sk["equation"])

    sheet("3_Multiple_LR", "3a. Multiple linear regression coefficients",
          res.multiple["table"], note=res.multiple["equation"])
    sheet("3b_VIF", "3b. Variance inflation factors", res.multiple["vif"],
          note="VIF < 5 is comfortable, 5-10 moderate, above 10 severe.")
    sheet("3c_Correlation", "3c. Predictor correlation matrix",
          res.multiple["correlation"].reset_index().rename(columns={"index": ""}))

    sheet("4_Diagnostics", "4. Regression diagnostics", res.diagnostics["summary_table"],
          note="Durbin-Watson near 2 means no autocorrelation. BP and JB use alpha = 0.05.")
    sheet("4b_Residuals", "4b. Fitted values and residuals",
          pd.DataFrame({"fitted": res.diagnostics["fitted"].values,
                        "residual": res.diagnostics["residuals"].values}))

    sheet("5_Polynomial", "5. Polynomial regression - degree comparison",
          res.polynomial["table"],
          note="In-sample fit always improves with degree; compare adjusted R2 and AIC.")

    sheet("6_Logistic", "6. Logistic regression for High_Cases", res.logistic["table"],
          note=res.logistic["note"] + " | pseudo R2 = %.4f" % res.logistic["pseudo_r2"])
    ws = wb["6_Logistic"]
    ws.cell(row=res.logistic["table"].shape[0] + 6, column=1,
            value="Full sm.Logit summary").font = Font(bold=True, color=NAVY)
    for i, line in enumerate(str(res.logistic["summary"]).splitlines(), 1):
        ws.cell(row=res.logistic["table"].shape[0] + 6 + i, column=1, value=line).font = Font(
            name="Consolas", size=8)

    cm = res.classification["confusion_matrix"].reset_index().rename(columns={"index": "actual"})
    sheet("7_Confusion", "7a. Confusion matrix at the 0.5 cut-off", cm)
    sheet("7b_Metrics", "7b. Classification metrics", res.classification["metrics"])
    sheet("7c_ROC", "7c. ROC curve coordinates",
          pd.DataFrame({"false_positive_rate": res.classification["fpr"],
                        "true_positive_rate": res.classification["tpr"],
                        "threshold": res.classification["thresholds"]}),
          note="AUC = %.4f" % res.classification["auc"])

    sheet("8_CrossValidation", "8. K-fold cross-validation with Ridge and Lasso",
          res.cv["table"], note=f"{res.cv['folds']}-fold KFold, shuffled with seed {SEED}.")
    sheet("8b_Coefficients", "8b. Standardised coefficients", res.cv["coefficients"])

    wb.save(path)
    print("  wrote", path)
    return path


if __name__ == "__main__":
    res = run_all()
    export_report(res)
